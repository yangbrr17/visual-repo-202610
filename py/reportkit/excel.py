"""Excel export (``build(..., excel=True)``): every dataset as one styled worksheet.

Each sheet: dataset title, source / notes / unit / as-of rows, a bold header with the field titles,
the data with Excel number formats taken from the fields' format codes (the same Excel-style codes the
report uses), frozen header, autofilter and column widths. openpyxl is imported lazily — it is only
needed when ``excel=True``.

Output is deterministic: workbook timestamps are the build date and zip entries carry a fixed time,
so rebuilding with the same input gives a byte-identical file.
"""
from __future__ import annotations

import datetime as _dt
import io
import json
import re
import zipfile
from typing import TYPE_CHECKING, Any

from ._util import as_list, colon, label, list_sep, source_text
from .data import is_iso_date
from .errors import BuildError
from .template import fallback_text

if TYPE_CHECKING:  # pragma: no cover
    from .report import Report

__all__ = ["excel_format", "is_date_format", "write_workbook", "workbook_bytes", "is_reportkit_xlsx", "PRESETS"]

# kit/core/rk.format.js presets (named formats usable instead of codes)
PRESETS = {
    "int": "#,##0", "num": "#,##0.0", "num1": "#,##0.0", "num2": "#,##0.00",
    "pct": "0%", "pct0": "0%", "pct1": "0.0%", "pct2": "0.00%",
    "delta_pct": "+0.0%;-0.0%;0.0%", "delta": "+#,##0.0;-#,##0.0;0.0",
    "mult": "0.0x", "multiple": "0.0x", "mult2": "0.00x",
    "usd": "$#,##0", "usd2": "$#,##0.00", "usd_m": '$#,##0"m"', "usd_bn": '$#,##0.0"bn"', "eur_m": '€#,##0"m"',
    "bps": '#,##0" bps"', "acct": "#,##0_);(#,##0)", "acct1": "#,##0.0_);(#,##0.0)",
    "date": "dd mmm yyyy", "month": "mmm yy", "year": "yyyy", "auto": "General", "text": "@",
}

# characters Excel displays literally without quoting (number sections)
_NUM_LITERAL_OK = set("$-+/():!^&'~{}<>= ")
_NUM_CODE = set("0#?.,%")
_DATE_CODE = set("dmyhsDMYHS")
_DATE_LITERAL_OK = set("/-:., ")
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MAX_CELL = 32767
_DATE_TYPES = ("date",)
_CREATOR = "ReportKit"


def _split_sections(code: str) -> list[str]:
    out, cur, in_q, i = [], "", False, 0
    while i < len(code):
        c = code[i]
        if c == '"':
            in_q = not in_q
        if c == "\\" and not in_q:
            cur += code[i:i + 2]
            i += 2
            continue
        if c == ";" and not in_q:
            out.append(cur)
            cur = ""
            i += 1
            continue
        cur += c
        i += 1
    out.append(cur)
    return out


def _is_date_section(sec: str) -> bool:
    bare = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", sec)
    return not re.search(r"[0#?]", bare) and bool(re.search(r"[dmyhsDMYHS]", bare))


def _convert_section(sec: str) -> str:
    if sec.strip().lower() == "general":
        return "General"
    is_date = _is_date_section(sec)
    out: list[str] = []
    lit: list[str] = []

    def flush() -> None:
        if lit:
            out.append('"' + "".join(lit).replace('"', "") + '"')
            lit.clear()

    i, n = 0, len(sec)
    while i < n:
        c = sec[i]
        if c == '"':  # quoted literal: keep as is
            j = sec.find('"', i + 1)
            j = n if j < 0 else j
            flush()
            out.append(sec[i:j + 1] if j < n else sec[i:] + '"')
            i = j + 1
            continue
        if c == "\\" and i + 1 < n:
            flush()
            out.append(sec[i:i + 2])
            i += 2
            continue
        if c in "_*" and i + 1 < n:   # _x (space of width x) / *x (repeat fill)
            flush()
            out.append(sec[i:i + 2])
            i += 2
            continue
        if c == "[":
            j = sec.find("]", i)
            j = n - 1 if j < 0 else j
            flush()
            out.append(sec[i:j + 1])
            i = j + 1
            continue
        if c == "@":
            flush()
            out.append(c)
            i += 1
            continue
        if is_date:
            if sec[i:i + 5].upper() == "AM/PM":
                flush()
                out.append(sec[i:i + 5])
                i += 5
                continue
            if sec[i:i + 3].upper() == "A/P":
                flush()
                out.append(sec[i:i + 3])
                i += 3
                continue
            if c in _DATE_CODE or c in _DATE_LITERAL_OK or c in "0":
                flush()
                out.append(c)
            else:
                lit.append(c)
            i += 1
            continue
        if c in "Ee" and i + 1 < n and sec[i + 1] in "+-" and out:  # scientific: 0.00E+00
            flush()
            out.append(sec[i:i + 2])
            i += 2
            continue
        if c in _NUM_CODE or c in "123456789" or c in _NUM_LITERAL_OK:
            flush()
            out.append(c)
        else:
            lit.append(c)
        i += 1
    flush()
    return "".join(out)


def is_date_format(code: str | None) -> bool:
    """True when every section of an (Excel) format code is a date / time pattern."""
    return bool(code) and all(_is_date_section(s) for s in _split_sections(code) if s.strip())


def excel_format(spec: Any) -> str | None:
    """ReportKit format code / preset name -> a valid Excel number format (None = General).

    The codes are Excel's own syntax, so they are kept; only characters Excel would reject unquoted
    (the 'x' of '0.0x', '€', letters) are wrapped in quotes: '0.0x' -> '0.0"x"'."""
    if not spec or not isinstance(spec, str):
        return None
    code = PRESETS.get(spec, spec)
    if code == "General":
        return None
    return ";".join(_convert_section(s) for s in _split_sections(code))


# ---------------------------------------------------------------------------------------------
# cell values
# ---------------------------------------------------------------------------------------------
def _clean_text(s: str) -> str:
    s = _ILLEGAL_XML.sub("", s)
    return s if len(s) <= _MAX_CELL else s[:_MAX_CELL - 1] + "…"


def _cell_value(v: Any, ftype: str) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (list, dict)):
        return _clean_text(json.dumps(v, ensure_ascii=False))
    s = str(v)
    if ftype in _DATE_TYPES and is_iso_date(s):
        try:
            if len(s) == 10:
                return _dt.date.fromisoformat(s)
            d = _dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
            return d.replace(tzinfo=None)
        except ValueError:
            pass
    return _clean_text(s)


def _display_len(v: Any) -> int:
    if v is None:
        return 0
    if isinstance(v, float):
        return min(len(f"{v:,.2f}"), 18)
    if isinstance(v, int) and not isinstance(v, bool):
        return len(f"{v:,}")
    if isinstance(v, (_dt.date, _dt.datetime)):
        return 11
    return max((len(line) for line in str(v).splitlines()), default=0)


def _wrapped_lines(text: str, width: float) -> int:
    """Lines a header needs when word-wrapped at *width* characters (greedy, like Excel)."""
    lines, cur = 1, 0
    for word in text.split():
        if cur and cur + 1 + len(word) > width:
            lines, cur = lines + 1, len(word)
        else:
            cur = cur + (1 if cur else 0) + len(word)
    return lines


def _sheet_names(report: "Report") -> dict[str, str]:
    used: set[str] = set()
    out: dict[str, str] = {}
    for did in report.datasets:
        base = re.sub(r"[\[\]:*?/\\]", "_", did).strip("'") or "data"
        if base.lower() == "history":   # reserved by Excel
            base += "_"
        base = base[:31]
        name, n = base, 2
        while name.lower() in used:
            suffix = f"~{n}"
            name = base[:31 - len(suffix)] + suffix
            n += 1
        used.add(name.lower())
        out[did] = name
    return out


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as e:
        raise BuildError("excel=True needs openpyxl: pip install openpyxl  (or pip install 'reportkit[excel]')") from e
    return openpyxl


def workbook_bytes(report: "Report", *, build_date: str | None = None) -> bytes:
    """The .xlsx file content for *report* (one sheet per dataset)."""
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    loc = report.locale
    sep = colon(loc)
    stamp = _stamp(build_date or report.date)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.creator = _CREATOR
    wb.properties.lastModifiedBy = _CREATOR
    wb.properties.title = fallback_text(report.title)
    if report.subtitle:
        wb.properties.subject = fallback_text(report.subtitle)
    wb.properties.created = stamp
    wb.properties.modified = stamp

    ink = "1B2A41"
    title_font = Font(name="Calibri", size=14, bold=True, color=ink)
    meta_font = Font(name="Calibri", size=9, italic=True, color="5A6472")
    head_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=ink)
    head_border = Border(bottom=Side(style="thin", color=ink))
    head_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    names = _sheet_names(report)
    for did, ds in report.datasets.items():
        ws = wb.create_sheet(names[did])
        fields = (ds.get("schema") or {}).get("fields", [])
        meta = ds.get("meta") or {}
        row = 1
        ws.cell(row=row, column=1, value=_clean_text(str(ds.get("title") or did))).font = title_font
        row += 1
        lines: list[str] = []
        if as_list(meta.get("source")):
            lines.append(source_text(meta.get("source"), loc))
        notes = [str(n) for n in as_list(meta.get("notes"))]
        if len(notes) == 1:
            lines.append(label("note", loc) + sep + notes[0])
        elif notes:   # one numbered note per row (long notes stay readable)
            lines.append(label("notes", loc) + sep + f"(1) {notes[0]}")
            lines += [f"({i}) {n}" for i, n in enumerate(notes[1:], 2)]
        extra = []
        if meta.get("unit"):
            extra.append(label("unit", loc) + sep + str(meta["unit"]))
        if meta.get("asOf"):
            extra.append(label("as_of", loc) + sep + str(meta["asOf"]))
        if extra:
            lines.append("  ·  ".join(extra))
        computed = [c for c in ds.get("computed") or [] if c.get("name") and c.get("formula")]
        if computed:
            lines.append(label("computed", loc) + sep + list_sep(loc).join(
                f"{c.get('title') or c['name']} = {c['formula']}" for c in computed))
        if report.confidentiality:
            lines.append(str(report.confidentiality))
        for text in lines:
            ws.cell(row=row, column=1, value=_clean_text(text)).font = meta_font
            row += 1
        row += 1  # spacer

        header_row = row
        widths: list[int] = []
        for ci, f in enumerate(fields, 1):
            title = str(f.get("title") or f["name"])
            unit = f.get("unit")
            if unit and str(unit).lower() not in title.lower():
                title = f"{title} ({unit})"
            c = ws.cell(row=header_row, column=ci, value=_clean_text(title))
            c.font, c.fill, c.alignment, c.border = head_font, head_fill, head_align, head_border
            widths.append(min(max(len(w) for w in title.split()) if title.split() else 4, 30))

        fmts = [excel_format(f.get("format")) for f in fields]
        date_fmt = [is_date_format(f) for f in fmts]
        date_default = "dd mmm yyyy"
        for ri, rec in enumerate(ds.get("data") or [], header_row + 1):
            for ci, f in enumerate(fields, 1):
                ftype = f.get("type") or "string"
                v = _cell_value(rec.get(f["name"]) if isinstance(rec, dict) else None, ftype)
                c = ws.cell(row=ri, column=ci, value=v)
                if isinstance(v, str):
                    c.data_type = "s"   # never a formula: "=HYPERLINK(...)" stays text
                fmt = fmts[ci - 1]
                if isinstance(v, (_dt.date, _dt.datetime)):
                    c.number_format = fmt or date_default
                elif fmt and isinstance(v, (int, float)) and not isinstance(v, bool) and not date_fmt[ci - 1]:
                    c.number_format = fmt   # (a date code on a plain number, e.g. 'yyyy' on 2025, stays General)
                widths[ci - 1] = max(widths[ci - 1], min(_display_len(v) + 1, 50))
        last_row = header_row + max(len(ds.get("data") or []), 0)
        if fields:
            last_col = get_column_letter(len(fields))
            ws.auto_filter.ref = f"A{header_row}:{last_col}{max(last_row, header_row)}"
            ws.freeze_panes = f"A{header_row + 1}"
            lines_needed = 1
            for ci, w in enumerate(widths, 1):
                width = max(8, min(w + 2, 52))
                ws.column_dimensions[get_column_letter(ci)].width = width
                head = str(ws.cell(row=header_row, column=ci).value or "")
                lines_needed = max(lines_needed, _wrapped_lines(head, width - 1))
            ws.row_dimensions[header_row].height = 14 * min(lines_needed, 4) + 4
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = f"{header_row}:{header_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = _footer_text(fallback_text(report.title))
        if report.confidentiality:
            ws.oddFooter.center.text = _footer_text(str(report.confidentiality))
        ws.oddFooter.right.text = "&P / &N"

    if not wb.sheetnames:
        wb.create_sheet("data")
    buf = io.BytesIO()
    wb.save(buf)
    return _normalise_zip(buf.getvalue(), stamp)


def _footer_text(s: str) -> str:
    """Header / footer text: '&' starts a control code there ('&C' = centre), so literal '&' is '&&'."""
    return _clean_text(s)[:120].replace("&", "&&")


def _stamp(date: Any) -> _dt.datetime:
    try:
        return _dt.datetime.fromisoformat(str(date)[:10])
    except (TypeError, ValueError):
        return _dt.datetime(2000, 1, 1)


def _normalise_zip(data: bytes, stamp: _dt.datetime) -> bytes:
    """Rewrite the archive with a fixed entry time — and the core 'modified' property, which openpyxl
    always sets to the current time on save — so identical input -> identical bytes."""
    dt = (max(stamp.year, 1980), stamp.month, stamp.day, 0, 0, 0)
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ")
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              lambda m: m.group(1) + iso.encode() + m.group(2), body)
            zi = zipfile.ZipInfo(info.filename, date_time=dt)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, body)
    return out.getvalue()


def write_workbook(report: "Report", path, *, build_date: str | None = None) -> None:
    from pathlib import Path
    data = workbook_bytes(report, build_date=build_date)
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)


def is_reportkit_xlsx(path) -> bool:
    """True for a workbook a previous ReportKit build wrote (safe to replace / clean)."""
    try:
        with zipfile.ZipFile(path) as z:
            core = z.read("docProps/core.xml").decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return False
    return re.search(rf"<dc:creator\b[^>]*>{_CREATOR}</dc:creator>", core) is not None
