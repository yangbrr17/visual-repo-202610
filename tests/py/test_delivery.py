"""Tests for the client-delivery features of the builder:

- I1 title templates (data-rk-template + fallback text, validation)
- I2 page(filters=False) -> data-rk-filters="off"
- I3 block option "choices" (validation, initial encoding)
- I5 Report(locale=, updown=) -> <html lang>, <meta name="rk:locale">, data-rk-updown
- I6 manifest optionDeps -> option-driven vendor deps; geo-world-50m vendor package
- B1 build(single_file=True)  (+ a Chromium check of the demo report as one file, if Playwright exists)
- B2 build(excel=True)
- B3 tools/vendor.py integrity hashes (--check)

Run:  cd <repo> && python3 -m unittest discover -s tests/py -v
Set REPORTKIT_SKIP_BROWSER=1 to skip the (slower) browser test.
"""
from __future__ import annotations

import contextlib
import datetime as dt
import html.parser
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "py"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from reportkit import BuildError, Kit, Report, ReportValidationError, SpecError  # noqa: E402
from reportkit import cli  # noqa: E402
from reportkit.build import inline_script_text, inline_style_text  # noqa: E402
from reportkit.template import fallback_text, parse_template  # noqa: E402
from test_builder import KitTestCase, _w, parse_bundle, read, tree  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAVE_OPENPYXL = False


def _load_tool(name: str):
    spec = importlib.util.spec_from_file_location(f"rk_tool_{name}", ROOT / "tools" / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class _ScriptCounter(html.parser.HTMLParser):
    """Counts elements the way a browser tokenises them (a stray '</script>' would end a script early)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tags: dict[str, int] = {}
        self.scripts: list[str] = []
        self._in_script = False

    def handle_starttag(self, tag, attrs):
        self.tags[tag] = self.tags.get(tag, 0) + 1
        if tag == "script":
            self._in_script = True
            self.scripts.append("")

    def handle_endtag(self, tag):
        if tag == "script":
            self._in_script = False

    def handle_data(self, data):
        if self._in_script:
            self.scripts[-1] += data


class DeliveryTestCase(KitTestCase):
    """Fake kit + a map component with optionDeps, fonts referenced from CSS, vendor CSS with an image."""

    def setUp(self):
        super().setUp()
        k = self.kit_root
        _w(k / "kit/styles/tokens.css",
           '/* TOKENS */ @font-face { font-family: "Inter"; src: url("../fonts/inter/inter.woff2") format("woff2"); }\n'
           ":root{--x:1}\n")
        _w(k / "kit/components/worldmap/manifest.json", json.dumps({
            "id": "worldmap", "name": "World map", "category": "geo", "deps": ["leaflet", "geo-a"],
            "channels": {"value": {"type": "number", "required": True}},
            "options": {"basemap": {"type": "string", "default": "110m"}, "heat": {"default": False}},
            "optionDeps": {"basemap": {"50m": ["geo-b"]}, "heat": {"true": ["heat-lib"]},
                           "layers": {"clusters": ["leaflet-markercluster"]}, "view.zoom": {"3": ["geo-b"]}},
        }))
        _w(k / "kit/components/worldmap/worldmap.js", "/* WORLDMAP-JS \u2028 */\n")
        vj = json.loads(read(k / "vendor/vendor.json"))
        vj["packages"] += [{"id": "geo-a", "local": True, "files": {"a.js": "a.js", "LICENSE": "LICENSE"}},
                           {"id": "geo-b", "local": True, "files": {"b.js": "b.js", "LICENSE": "LICENSE"}},
                           {"id": "heat-lib", "files": {"heat.js": "heat.js"}, "requires": ["leaflet"]}]
        _w(k / "vendor/vendor.json", json.dumps(vj))
        _w(k / "vendor/geo-a/a.js", "/* GEO-A */\n")
        _w(k / "vendor/geo-a/LICENSE", "ISC\n")
        _w(k / "vendor/geo-b/b.js", "/* GEO-B */\n")
        _w(k / "vendor/geo-b/LICENSE", "ISC\n")
        _w(k / "vendor/heat-lib/heat.js", "/* HEAT */\n")
        _w(k / "vendor/leaflet/leaflet.css", "/* LEAFLET-CSS */ .x{background:url(images/marker.png)} .v{behavior:url(#default#VML)}\n")

    def quiet_build(self, r: Report, out=None, **kw):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return r.build(out or self.out, **kw)


# ---------------------------------------------------------------------------------------------
# I1 title templates
# ---------------------------------------------------------------------------------------------
class TemplateTests(DeliveryTestCase):
    def test_parse_placeholders(self):
        ph = parse_template("A {{peers: median(ev_ebitda) | 0.0x}} B {{@wacc | 0.0%}} C {{@wacc}}")
        self.assertEqual([(p.dataset, p.expr, p.format) for p in ph],
                         [("peers", "median(ev_ebitda)", "0.0x"), (None, "@wacc", "0.0%"), (None, "@wacc", None)])
        # '||' is the or-operator, never the separator; '|' inside quotes is not a separator either
        ph = parse_template('{{d: a > 1 || b > 2}} {{d.v2: count(sector == "A|B") | #,##0}} {{x ? 1 : 2}}')
        self.assertEqual([(p.dataset, p.expr, p.format) for p in ph],
                         [("d", "a > 1 || b > 2", None), ("d.v2", 'count(sector == "A|B")', "#,##0"),
                          (None, "x ? 1 : 2", None)])
        self.assertEqual(parse_template('{{d: sum(ev) | $#,##0.0,,"bn"}}')[0].format, '$#,##0.0,,"bn"')
        self.assertEqual(fallback_text("Trades at {{d: x}} vs {{d: y | 0.0x}}"), "Trades at \u2026 vs \u2026")
        self.assertEqual(fallback_text("No template"), "No template")

    def test_render_template_attributes_and_fallback(self):
        r = self.make_report()
        p = r.page('Peers trade at {{peers: median(ev_ebitda) | 0.0x}} ("{{peers: count()}}" names)',
                   kicker="{{@adj | 0.00x}} adjusted", subtitle='Max EV {{peers: max(ev) | $#,##0.0,,"bn"}}')
        p.block("chart", data="peers", encode={"x": "company", "y": "ev"},
                title="EV of {{peers: company}}", subtitle="{{peers: sum(ev) / sum(ebitda) | 0.0x}} aggregate")
        r.cover("Atlas at {{peers: median(ev) | $#,##0}}")
        r.validate()
        html = r.html()
        self.assertIn('<h2 class="rk-page__title" data-rk-template="Peers trade at {{peers: median(ev_ebitda) | 0.0x}} '
                      '(&quot;{{peers: count()}}&quot; names)">Peers trade at \u2026 (&quot;\u2026&quot; names)</h2>', html)
        self.assertIn('<p class="rk-page__kicker" data-rk-template="{{@adj | 0.00x}} adjusted">\u2026 adjusted</p>', html)
        self.assertIn('<p class="rk-page__subtitle" data-rk-template="Max EV {{peers: max(ev) | $#,##0.0,,&quot;bn&quot;}}">'
                      'Max EV \u2026</p>', html)
        self.assertIn('<h3 class="rk-block__title" data-rk-template="EV of {{peers: company}}">EV of \u2026</h3>', html)
        self.assertIn('<p class="rk-block__subtitle" data-rk-template="{{peers: sum(ev) / sum(ebitda) | 0.0x}} aggregate">'
                      '\u2026 aggregate</p>', html)
        self.assertIn('<h1 class="rk-cover__title" data-rk-template="Atlas at {{peers: median(ev) | $#,##0}}">Atlas at \u2026</h1>',
                      html)
        # ids and the nav label come from the fallback text (no braces / expressions in ids)
        self.assertIn('id="adjusted" data-rk-page="\u2026 adjusted"', html)
        self.assertIn('id="adjusted-ev-of"', html)
        # plain titles are unchanged (no attribute)
        self.assertIn('<h2 class="rk-page__title">Peers trade at a median 9.5x EBITDA</h2>', html)
        self.assertEqual(html.count("data-rk-template="), 6)

    def test_template_validation(self):
        r = self.make_report()
        p = r.page("Bad {{peerz: median(ev)}}", kicker="{{peers: medain_ev | 0.0x}}",
                   subtitle="{{@wac | 0.0%}} and {{ev / 2}}")
        p.block("chart", data="peers", encode={"x": "company", "y": "ev"}, title="{{peers: ev # 2}}",
                subtitle="Unclosed {{peers: ev")
        r.page("{{peers: }} empty", subtitle="{{peers: ev | }} ok: empty format = inferred")
        r.page("No data-bound block: {{median(ev)}}")
        errors, _ = r.check()
        text = "\n".join(errors)
        self.assertIn("unknown dataset 'peerz' — did you mean 'peers'?", text)
        self.assertIn("field 'medain_ev' is not in dataset 'peers'", text)
        self.assertIn("uses @wac but no param 'wac' is defined", text)
        self.assertIn("'{{median(ev)}}' references field(s) ev but names no dataset", text)
        self.assertNotIn("'{{ev / 2}}'", text)   # page 4 has a data-bound block: its dataset is the default
        self.assertIn("expression 'ev # 2' is invalid (Unexpected character \"#\")", text)
        self.assertIn("unbalanced '{{' / '}}'", text)
        self.assertIn("empty expression", text)
        self.assertNotIn("ok: empty format", text)
        self.assertTrue(any("subtitle template" in e and "Bad" in e for e in errors), errors)
        self.assertTrue(any("block" in e and "› title template" in e for e in errors), errors)
        self.assertTrue(any("block" in e and "› subtitle template" in e for e in errors), errors)

    def test_template_accepts_titles_computed_fields_and_params(self):
        r = self.make_report()
        # field title ("EV"), computed field (ev_ebitda), [bracket name], @param, aggregate with filter arg
        r.page("{{peers: median([EV]) | $#,##0}} / {{peers: avg(ev_ebitda) * @adj | 0.0x}} / {{@adj}}")
        # no prefix: a block title defaults to the block's dataset, a page title to its first data-bound block
        p = r.page("Median {{median(ev) | $#,##0}}", subtitle="{{@adj * 2}} params need no dataset")
        p.block("note", title="Commentary")
        p.block("chart", data="peers", encode={"x": "company", "y": "ev"}, title="Top: {{max(ev)}}")
        errors, _ = r.check()
        self.assertEqual(errors, [])
        p.block("chart", data="peers", encode={"x": "company", "y": "ev"}, title="{{max(evv)}}")
        self.assertIn("field 'evv' is not in dataset 'peers' — did you mean 'ev'?", "\n".join(r.check()[0]))
        # unknown block dataset: reported once on the block, the template check does not crash
        r.page("P").block("chart", data="nope", encode={"x": "company"}, title="{{max(ev)}}")
        errors = "\n".join(r.check()[0])
        self.assertIn("unknown dataset 'nope'", errors)

    def test_template_spec_round_trip(self):
        r = self.make_report()
        r.page("At {{peers: median(ev) | $#,##0}}", filters=False)
        r2 = Report.from_spec(json.loads(json.dumps(r.to_spec())))
        self.assertEqual(r.html(), r2.html())


# ---------------------------------------------------------------------------------------------
# I2 page filters / I5 locale + updown / I3 choices
# ---------------------------------------------------------------------------------------------
class PageOptionTests(DeliveryTestCase):
    def test_filters_off(self):
        r = self.make_report()
        r.page("Unfiltered market overview", kicker="Market", filters=False).block(
            "chart", data="peers", encode={"x": "company", "y": "ev"})
        r.validate()
        html = r.html()
        self.assertIn('<section class="rk-page" id="market" data-rk-page="Market" data-rk-filters="off">', html)
        self.assertEqual(html.count("data-rk-filters"), 1)   # default pages carry no attribute
        spec = r.to_spec()
        self.assertIs(spec["pages"][-1]["filters"], False)
        self.assertNotIn("filters", spec["pages"][-2])
        self.assertIn('data-rk-filters="off"', Report.from_spec(spec).html())
        r.page("X", filters="no")
        self.assertTrue(any("filters must be True or False" in e for e in r.check()[0]))

    def test_locale_and_updown(self):
        r = Report("项目", date="2026-10-14", locale="zh-CN", updown="cn")
        r.page("标题").block("note")
        html = r.html()
        self.assertTrue(html.startswith('<!DOCTYPE html>\n<html lang="zh-CN" data-rk-theme="navy" data-rk-updown="cn">'))
        self.assertIn('<meta name="rk:locale" content="zh-CN">', html)
        self.assertIn("打印 / PDF", html)
        self.assertEqual(r.manifest()["report"]["updown"], "cn")
        self.assertEqual(Report.from_spec(r.to_spec()).html(), html)
        en = Report("T")
        en.page("P").block("note")
        h2 = en.html()
        self.assertIn('<meta name="rk:locale" content="en-US">', h2)
        self.assertNotIn("data-rk-updown", h2)
        self.assertIn('data-rk-updown="us"', Report("T", updown="us").html())
        with self.assertRaises(ValueError):
            Report("T", updown="red")
        with self.assertRaises(SpecError):
            Report.from_spec({"report": {"title": "T", "updown": "eu"}, "pages": []})
        w = Report("T", locale="en_GB")
        w.page("P").block("note")
        self.assertTrue(any("BCP 47" in x for x in w.check()[1]))

    def test_choices(self):
        r = self.make_report()
        p = r.page("Pick a metric")
        b = p.block("chart", data="peers", encode={"x": "company"},
                    options={"choices": {"y": ["*computed", "ev", "ebitda"], "x": ["company", "sector"]}})
        self.assertEqual(r.check()[0], [])   # required channel y is satisfied by its choices
        html = r.html()
        opts = json.loads(html.split(f'id="{b.id}">')[1].split("data-rk-options>")[1].split("</script>")[0])
        self.assertEqual(opts["encode"], {"x": "company", "y": "ev"})        # first concrete choice
        self.assertEqual(b.options["encode"], {"x": "company"})              # the spec itself is untouched
        bad = r.page("Bad").block("chart", data="peers", encode={"x": "company", "y": "ev"},
                                  options={"choices": {"y": ["evv", "*all"], "x": []}})
        errors = "\n".join(r.check()[0])
        self.assertIn("options.choices.y -> field 'evv' is not in dataset 'peers' — did you mean 'ev'?", errors)
        self.assertIn("unknown wildcard '*all'", errors)
        self.assertIn("options.choices.x must be a non-empty list", errors)
        self.assertIsNotNone(bad)


# ---------------------------------------------------------------------------------------------
# I6 optionDeps
# ---------------------------------------------------------------------------------------------
class OptionDepsTests(DeliveryTestCase):
    def report(self, **options) -> Report:
        r = Report("Maps", date="2026-10-14")
        r.dataset("c", [{"iso": "840", "value": 1}])
        r.page("Map").block("worldmap", data="c", encode={"value": "value"}, options=options or None)
        return r

    def test_option_deps_resolution(self):
        m = Kit.locate().component("worldmap")
        self.assertEqual(Kit.option_deps(m, {}), [])
        self.assertEqual(Kit.option_deps(m, {"basemap": "50m"}), ["geo-b"])
        self.assertEqual(Kit.option_deps(m, {"heat": True}), ["heat-lib"])
        self.assertEqual(Kit.option_deps(m, {"heat": False}), [])
        self.assertEqual(Kit.option_deps(m, {"layers": ["labels", "clusters"]}), ["leaflet-markercluster"])
        self.assertEqual(Kit.option_deps(m, {"view": {"zoom": 3.0}}), ["geo-b"])
        self.assertEqual(Kit.option_deps(m, {"basemap": "50m", "view": {"zoom": 3}}), ["geo-b"])
        self.assertEqual(Kit.option_dep_ids(m), ["geo-b", "heat-lib", "leaflet-markercluster"])

    def test_build_ships_option_deps_only_when_used(self):
        out = self.quiet_build(self.report(), self.tmp / "plain")
        html = read(out / "index.html")
        self.assertIn('src="assets/vendor/geo-a/a.js"', html)
        self.assertNotIn("geo-b", html)
        self.assertFalse((out / "assets/vendor/geo-b").exists())
        out = self.quiet_build(self.report(basemap="50m", heat=True, layers=["clusters"]), self.tmp / "hi")
        html = read(out / "index.html")
        for src in ("assets/vendor/leaflet/leaflet.js", "assets/vendor/leaflet-markercluster/leaflet.markercluster.js",
                    "assets/vendor/geo-a/a.js", "assets/vendor/geo-b/b.js", "assets/vendor/heat-lib/heat.js"):
            self.assertIn(f'<script src="{src}"></script>', html)
        # vendor.json order: leaflet before its plugins, scripts before reportkit.js
        order = [html.index(f'<script src="assets/{s}') for s in (
            "vendor/leaflet/leaflet.js", "vendor/leaflet-markercluster/", "vendor/geo-a/a.js", "vendor/geo-b/b.js",
            "vendor/heat-lib/heat.js", "js/reportkit.js")]
        self.assertEqual(order, sorted(order))
        self.assertTrue((out / "assets/vendor/geo-b/LICENSE").is_file())

    def test_unknown_option_dep_vendor_is_a_kit_error(self):
        mpath = self.kit_root / "kit/components/worldmap/manifest.json"
        m = json.loads(read(mpath))
        m["optionDeps"]["basemap"]["10m"] = ["geo-ten"]
        _w(mpath, json.dumps(m))
        errors = "\n".join(self.report().check()[0])
        self.assertIn("optionDeps names unknown vendor package 'geo-ten'", errors)

    def test_missing_option_dep_files_reported(self):
        (self.kit_root / "vendor/geo-b/b.js").unlink()
        self.assertEqual(self.report().check()[0], [])          # not used -> not checked
        errors = "\n".join(self.report(basemap="50m").check()[0])
        self.assertIn("vendor/geo-b/ is incomplete (missing b.js)", errors)


# ---------------------------------------------------------------------------------------------
# B1 single file
# ---------------------------------------------------------------------------------------------
class SingleFileTests(DeliveryTestCase):
    def make(self) -> Report:
        r = self.make_report()
        r.dataset("evil", [{"k": "</script><script>alert(1)</script>", "c": "<!-- x --> \u2028 </SCRIPT >", "value": 1}])
        r.page("Map", kicker="Geo").block("worldmap", data="evil", encode={"value": "value"},
                                         options={"basemap": "50m", "layers": ["clusters"]})
        return r

    def test_one_self_contained_file(self):
        path = self.quiet_build(self.make(), self.tmp / "deliver", single_file=True)
        self.assertEqual(path, (self.tmp / "deliver/project-atlas.html").resolve())
        self.assertEqual(sorted(p.name for p in path.parent.iterdir()), ["project-atlas.html"])
        html = read(path)
        # nothing is referenced from disk: no <link href>, no <script src>, CSS url()s are data: URIs
        self.assertNotIn("<link", html)
        self.assertNotIn("<script src", html)
        self.assertIn('url("data:font/woff2;base64,Rk9OVA==")', html)       # kit/fonts/inter/inter.woff2 = "FONT"
        self.assertIn('url("data:image/png;base64,UE5H")', html)            # vendor/leaflet/images/marker.png = "PNG"
        self.assertIn("url(#default#VML)", html)                            # fragment refs left alone
        for marker in ("CORE-1", "CORE-2", "TOKENS", "NAVY", "CHART-CSS", "CHART-JS", "FILTER-JS", "NOTE-JS",
                       "ECHARTS", "LEAFLET", "LEAFLET-CSS", "MC-CSS", "GEO-A", "GEO-B", "WORLDMAP-JS"):
            self.assertIn(marker, html, marker)
        self.assertNotIn("UNUSED", html)
        # order: vendor CSS before reportkit.css; vendor JS -> core -> components -> data
        order = [html.index(m) for m in ("LEAFLET-CSS", "/* TOKENS */", "/* ECHARTS */", "/* CORE-1 */",
                                         "/* CHART-JS */", "RK.data.preload(")]
        self.assertEqual(order, sorted(order))
        self.assertIn('<meta name="generator" content="ReportKit">', html)

    def test_inline_escaping_keeps_scripts_intact(self):
        path = self.quiet_build(self.make(), self.tmp / "x.html", single_file=True)
        self.assertEqual(path, (self.tmp / "x.html").resolve())
        html = read(path)
        self.assertNotIn("</script><script>alert(1)", html)
        self.assertNotIn("<!-- x -->", html)
        self.assertNotIn("\u2028", html)
        self.assertIn("WORLDMAP-JS \\u2028 */", html)   # raw U+2028 in JS -> escape
        parser = _ScriptCounter()
        parser.feed(html)
        # one <script> per inlined file (+ one per options JSON): no script was cut short by '</script'
        n_files = html.count("<script data-rk-inline=")
        self.assertEqual(n_files, 11)   # 5 vendor js + reportkit.js + 4 components + bundle
        self.assertEqual(parser.tags["script"], n_files + html.count("data-rk-options>"))
        bundle = next(s for s in parser.scripts if s.lstrip().startswith("RK.data.preload("))
        data = parse_bundle(bundle.strip())
        self.assertEqual(data["datasets"]["evil"]["data"][0]["k"], "</script><script>alert(1)</script>")
        self.assertEqual(data["datasets"]["evil"]["data"][0]["c"], "<!-- x --> \u2028 </SCRIPT >")

    def test_escape_helpers(self):
        self.assertEqual(inline_script_text('a("</script>"); b("<!--"); c("</SCRIPT")'),
                         'a("\\u003c/script>"); b("\\u003c!--"); c("\\u003c/SCRIPT")')
        self.assertEqual(inline_script_text("x = '\u2028\u2029'"), "x = '\\u2028\\u2029'")
        self.assertEqual(inline_script_text("a < b && c </ d"), "a < b && c </ d")
        self.assertEqual(inline_style_text('.a::after{content:"</style>"}'), '.a::after{content:"<\\/style>"}')

    def test_deterministic_overwrite_rules_and_modes(self):
        a = self.quiet_build(self.make(), self.tmp / "a", single_file=True)
        b = self.quiet_build(self.make(), self.tmp / "b", single_file=True)
        self.assertEqual(read(a), read(b))
        self.quiet_build(self.make(), self.tmp / "a", single_file=True)   # our own file: overwritten
        self.assertEqual(read(a), read(b))
        foreign = self.tmp / "c" / "project-atlas.html"
        foreign.parent.mkdir()
        foreign.write_text("<html>someone else's page</html>")
        with self.assertRaises(BuildError):
            self.quiet_build(self.make(), self.tmp / "c", single_file=True)
        self.assertIn("someone else", read(foreign))
        self.quiet_build(self.make(), self.tmp / "c", single_file=True, clean=False)
        self.assertIn("ReportKit", read(foreign))
        with self.assertRaises(BuildError):
            self.quiet_build(self.make(), self.tmp / "d", single_file=True, data_source="http")
        api = self.quiet_build(self.make(), self.tmp / "e", single_file=True, data_source="api",
                               api="https://example.test/api/r/1/")
        html = read(api)
        self.assertNotIn("RK.data.preload(", html)
        self.assertIn('<meta name="rk:api" content="https://example.test/api/r/1/">', html)

    def test_unresolvable_css_url_is_an_error(self):
        _w(self.kit_root / "kit/styles/tokens.css", '@font-face { src: url("../fonts/nope.woff2"); }\n')
        with self.assertRaises(BuildError) as cm:
            self.quiet_build(self.make_report(), self.tmp / "f", single_file=True)
        self.assertIn("nope.woff2", str(cm.exception))


# ---------------------------------------------------------------------------------------------
# B2 Excel
# ---------------------------------------------------------------------------------------------
class ExcelFormatTests(unittest.TestCase):
    def test_format_codes_map_to_valid_excel(self):
        from reportkit.excel import excel_format, is_date_format
        cases = {
            "0.0x": '0.0"x"', "0.00x": '0.00"x"', "mult": '0.0"x"', "0.0%": "0.0%", "pct1": "0.0%",
            "$#,##0": "$#,##0", '$#,##0.0,,"m"': '$#,##0.0,,"m"', "#,##0;(#,##0)": "#,##0;(#,##0)",
            "+0.0%;-0.0%;0.0%": "+0.0%;-0.0%;0.0%", "dd mmm yyyy": "dd mmm yyyy", "mmm yy": "mmm yy",
            "eur_m": '"€"#,##0"m"', "acct": "#,##0_);(#,##0)", '#,##0" bps"': '#,##0" bps"',
            "0.00E+00": "0.00E+00", "[Red]#,##0": "[Red]#,##0", "#,##0 USD": '#,##0 "USD"', "@": "@",
            "0.0\\x": "0.0\\x", "yyyy-mm-dd hh:mm": "yyyy-mm-dd hh:mm", "Q0": '"Q"0',
        }
        for src, want in cases.items():
            self.assertEqual(excel_format(src), want, src)
        self.assertIsNone(excel_format("auto"))
        self.assertIsNone(excel_format(None))
        self.assertTrue(is_date_format("dd mmm yyyy"))
        self.assertFalse(is_date_format('0.0"x"'))


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class ExcelTests(DeliveryTestCase):
    def make(self) -> Report:
        r = self.make_report()
        r.dataset("prices", [{"date": "2026-10-13", "close": 42.5, "mult": 9.64, "note": "=HYPERLINK(\"http://x\")",
                              "tags": ["a", "b"], "flag": True, "ctrl": "bell\x07"},
                             {"date": "2026-10-14", "close": None, "mult": 10.1, "note": "@SUM(1)", "tags": None,
                              "flag": False, "ctrl": "ok"}],
                  title="Share prices", source=["FactSet", "Bloomberg"], notes=["Close", "Adjusted"], unit="USD",
                  as_of="2026-10-14",
                  fields={"date": {"title": "Date", "format": "dd mmm yyyy"},
                          "close": {"title": "Close", "format": "$#,##0.00", "unit": "USD"},
                          "mult": {"title": "EV / EBITDA", "format": "mult"}})
        return r

    def test_workbook_contents(self):
        import openpyxl
        out = self.quiet_build(self.make(), excel=True)
        xlsx = out / "data/project-atlas.xlsx"
        self.assertTrue(xlsx.is_file())
        self.assertIn('<a class="rk-btn" href="data/project-atlas.xlsx" download data-rk-download="excel"',
                      read(out / "index.html"))
        wb = openpyxl.load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["peers", "prices"])
        self.assertEqual(wb.properties.creator, "ReportKit")
        ws = wb["prices"]
        col = [ws.cell(row=i, column=1).value for i in range(1, 9)]
        self.assertEqual(col[0], "Share prices")
        self.assertTrue(ws.cell(row=1, column=1).font.b)
        self.assertEqual(col[1], "Sources: FactSet; Bloomberg")
        self.assertEqual(col[2], "Notes: (1) Close")
        self.assertEqual(col[3], "(2) Adjusted")
        self.assertEqual(col[4], "Unit: USD  ·  As of: 2026-10-14")
        self.assertEqual(col[5], "Strictly private & confidential")
        hdr = 8
        self.assertEqual(col[7], "Date")
        self.assertEqual(ws.freeze_panes, f"A{hdr + 1}")
        self.assertEqual(ws.auto_filter.ref, f"A{hdr}:G{hdr + 2}")
        heads = [ws.cell(row=hdr, column=c).value for c in range(1, 8)]
        self.assertEqual(heads, ["Date", "Close (USD)", "EV / EBITDA", "note", "tags", "flag", "ctrl"])
        self.assertTrue(all(ws.cell(row=hdr, column=c).font.b for c in range(1, 8)))
        d, close, mult, note, tags, flag, ctrl = (ws.cell(row=hdr + 1, column=c) for c in range(1, 8))
        self.assertEqual(d.value, dt.datetime(2026, 10, 13))
        self.assertEqual(d.number_format, "dd mmm yyyy")
        self.assertEqual((close.value, close.number_format), (42.5, "$#,##0.00"))
        self.assertEqual((mult.value, mult.number_format), (9.64, '0.0"x"'))
        self.assertEqual((note.value, note.data_type), ('=HYPERLINK("http://x")', "s"))   # never a formula
        self.assertEqual(tags.value, '["a", "b"]')
        self.assertIs(flag.value, True)
        self.assertEqual(ctrl.value, "bell")
        self.assertIsNone(ws.cell(row=hdr + 2, column=2).value)
        self.assertGreater(ws.column_dimensions["D"].width, 20)
        peers = wb["peers"]
        self.assertIn("Calculated live in the report", "\n".join(str(c.value) for c in peers["A"] if c.value))

    def test_deterministic_and_clean(self):
        import time
        a = self.quiet_build(self.make(), self.tmp / "a", excel=True)
        time.sleep(1.1)   # openpyxl stamps 'modified' with the current second; the builder pins it
        b = self.quiet_build(self.make(), self.tmp / "b", excel=True)
        self.assertEqual(tree(a), tree(b))
        (a / "data/analyst.xlsx").write_bytes(b"not ours")
        r = self.make()
        r.title = "Project Borealis"
        self.quiet_build(r, a, excel=True)
        self.assertFalse((a / "data/project-atlas.xlsx").exists())     # stale export of ours removed
        self.assertTrue((a / "data/project-borealis.xlsx").exists())
        self.assertTrue((a / "data/analyst.xlsx").exists())            # foreign workbook kept
        self.quiet_build(r, a)                                          # excel=False: no workbook, no link
        self.assertFalse((a / "data/project-borealis.xlsx").exists())
        self.assertNotIn("data-rk-download", read(a / "index.html"))

    def test_single_file_with_excel(self):
        path = self.quiet_build(self.make(), self.tmp / "sf", single_file=True, excel=True)
        self.assertTrue(path.with_suffix(".xlsx").is_file())
        self.assertIn('href="project-atlas.xlsx" download', read(path))

    def test_cli_flags(self):
        spec = self.make().save_spec(self.tmp / "spec/report.json")
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
            code = cli.main(["build", str(spec), str(self.tmp / "cli"), "--single-file", "--excel", "--locale", "zh-CN"])
        self.assertEqual(code, 0)
        html = read(self.tmp / "cli/project-atlas.html")
        self.assertIn('<html lang="zh-CN"', html)
        self.assertIn('<meta name="rk:locale" content="zh-CN">', html)
        self.assertTrue((self.tmp / "cli/project-atlas.xlsx").is_file())
        self.assertIn("project-atlas.xlsx", buf.getvalue())


class ExcelMissingTests(DeliveryTestCase):
    def test_clear_error_without_openpyxl(self):
        saved = {k: v for k, v in sys.modules.items() if k == "openpyxl" or k.startswith("openpyxl.")}
        for k in saved:
            del sys.modules[k]
        sys.modules["openpyxl"] = None   # import openpyxl -> ImportError
        try:
            with self.assertRaises(BuildError) as cm:
                self.quiet_build(self.make_report(), excel=True)
            self.assertIn("pip install openpyxl", str(cm.exception))
            self.assertFalse(self.out.exists())              # failed before writing anything
            self.quiet_build(self.make_report())              # excel=False never needs it
        finally:
            del sys.modules["openpyxl"]
            sys.modules.update(saved)


# ---------------------------------------------------------------------------------------------
# B3 vendor integrity + I6 vendor package
# ---------------------------------------------------------------------------------------------
class VendorIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.vendor = _load_tool("vendor")
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_repo_vendor_tree_matches_recorded_hashes(self):
        manifest = self.vendor.load_manifest(ROOT)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(self.vendor.check(manifest, ROOT), 0)
        for p in manifest["packages"]:
            self.assertTrue(p.get("sha256"), p["id"])
            if not p.get("local"):
                self.assertTrue(p.get("integrity", "").startswith("sha512-"), p["id"])
        ids = [p["id"] for p in manifest["packages"]]
        geo = manifest["packages"][ids.index("geo-world-50m")]
        self.assertTrue(geo.get("local"))
        self.assertEqual(sorted(geo["sha256"]), ["LICENSE", "world-50m.js"])
        head = (ROOT / "vendor/geo-world-50m/world-50m.js").read_text(encoding="utf-8")
        self.assertIn("RK.geo.register('world-50m',d)", head[-200:])
        self.assertIn("geo-world-50m", (ROOT / "vendor/LICENSES.md").read_text(encoding="utf-8"))

    def test_check_detects_tampering_and_strays(self):
        v = self.vendor
        _w(self.tmp / "vendor/lib/lib.js", "/* LIB */\n")
        _w(self.tmp / "vendor/lib/LICENSE", "MIT\n")
        _w(self.tmp / "vendor/lib/img/a.png", "A")
        manifest = {"packages": [{"id": "lib", "npm": "lib", "version": "1.0.0", "license": "MIT", "repo": "-",
                                  "files": {"dist/lib.js": "lib.js", "LICENSE": "LICENSE", "dist/img/": "img/"}}]}
        quiet = contextlib.redirect_stdout(io.StringIO())
        with quiet:
            self.assertEqual(v.check(manifest, self.tmp), 1)          # no hashes recorded yet
        v.rehash(manifest, root=self.tmp)
        self.assertEqual(sorted(manifest["packages"][0]["sha256"]), ["LICENSE", "img/a.png", "lib.js"])
        v.save_manifest(manifest, self.tmp)
        self.assertEqual(v.load_manifest(self.tmp), manifest)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(v.check(manifest, self.tmp), 0)
        (self.tmp / "vendor/lib/lib.js").write_text("/* LIB */ steal()\n")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            self.assertEqual(v.check(manifest, self.tmp), 1)
        self.assertIn("sha256 mismatch: vendor/lib/lib.js", out.getvalue())
        v.rehash(manifest, root=self.tmp)
        _w(self.tmp / "vendor/lib/extra.js", "evil")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            self.assertEqual(v.check(manifest, self.tmp), 1)
        self.assertIn("unexpected file: vendor/lib/extra.js", out.getvalue())

    def test_vendor_cli_check(self):
        res = subprocess.run([sys.executable, str(ROOT / "tools/vendor.py"), "--check"], capture_output=True, text=True)
        self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
        self.assertIn("vendor check: OK", res.stdout)


# ---------------------------------------------------------------------------------------------
# B1 in a real browser: the demo pitchbook as one file, opened from file://
# ---------------------------------------------------------------------------------------------
def _playwright_available() -> bool:
    if os.environ.get("REPORTKIT_SKIP_BROWSER") or not shutil.which("node"):
        return False
    probe = subprocess.run(["node", "-e", "try{require('playwright')}catch(e){require(process.env.PLAYWRIGHT_PATH||"
                                          "'/opt/node-tools/node_modules/playwright')}"], capture_output=True)
    return probe.returncode == 0


@unittest.skipUnless(_playwright_available(), "node + playwright not available (or REPORTKIT_SKIP_BROWSER set)")
class SingleFileBrowserTests(unittest.TestCase):
    def test_demo_single_file_renders_offline(self):
        spec = importlib.util.spec_from_file_location("rk_build_demo", ROOT / "py/examples/build_demo.py")
        demo = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(demo)
        old = os.environ.pop("REPORTKIT_HOME", None)
        try:
            with tempfile.TemporaryDirectory() as tmp, warnings.catch_warnings():
                warnings.simplefilter("ignore")
                path = demo.build(Path(tmp) / "delivery", single_file=True)
                self.assertEqual(sorted(p.name for p in path.parent.iterdir()), ["project-atlas.html"])
                res = subprocess.run(["node", str(ROOT / "tests/py/single_file_check.cjs"), str(path)],
                                     capture_output=True, text=True, timeout=240)
        finally:
            if old is not None:
                os.environ["REPORTKIT_HOME"] = old
        self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
        info = json.loads(res.stdout.strip().splitlines()[-1])
        self.assertGreater(info["blocks"], 20)
        self.assertEqual(info["notRendered"], [])
        self.assertEqual(info["errors"], [])
        self.assertEqual(info["requests"], [])        # nothing but data: URIs
        self.assertEqual(info["externalRefs"], [])
        self.assertTrue(any(f.startswith("Inter") for f in info["fonts"]), info["fonts"])


if __name__ == "__main__":
    unittest.main()
